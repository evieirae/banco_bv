"""
Cross-check 06_prontidao_tabela*.csv (category=table_in) against src/RepMap.xlsx (réplicas Sybase / lineage).

Artefactos no inventário podem vir em notação SQL Server/Sybase a 3 partes:
  database.schema.table  — ex.: DBCOR.DBO.TBGRUPOCONDICAOPRECOPRCROCMRCL
onde o segmento do meio é o *owner* (schema) lógico — muitas vezes ``DBO`` (default). No RepMap a réplica
costuma aparecer como DBtarget=DBCOR + Tabela=TBGRUPO... (sem ``DBO`` no meio). Este script colapsa
owners conhecidos para comparar ``DBCOR`` + ``Tabela`` ao par ``DBtarget``/``DBsource`` + ``Tabela``.

Saída: indicativo YES/NO de réplica no RepMap, nível de match, e qual réplica (linha RepMap resumida).

Usage:
  python validate_table_in_vs_repmap.py [06_prontidao_tabela.csv] [src/RepMap.xlsx] [out.csv]
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path


def norm(s) -> str:
    return str(s or "").strip()


def norm_upper(s) -> str:
    return norm(s).upper()


# Owner/schema intermediário típico em 3-part naming (Sybase ASE / SQL Server); não existe como coluna separada no RepMap.
_SQL_OWNER_SKIP = frozenset(
    {
        "DBO",
        "SYS",
        "DBA",
        "GUEST",
        "INFORMATION_SCHEMA",
        "SYSIBM",
        "SYSCAT",
        "SYSPROC",
    }
)


def load_repmap_rows(xlsx_path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not row or all(v is None or str(v).strip() == "" for v in row[:8]):
            continue
        r = {
            "Source": norm(row[0]),
            "DBsource": norm(row[1]),
            "Target": norm(row[2]),
            "DBtarget": norm(row[3]),
            "Tipo": norm(row[4]),
            "Carga": norm(row[5]),
            "Tabela": norm(row[6]),
            "TabRep": norm(row[7]),
            "RepDef": norm(row[8]),
            "SubName": norm(row[9]) if len(row) > 9 else "",
            "_rownum": i + 2,
        }
        rows.append(r)
    return rows


def repmap_derived_keys(r: dict) -> set[str]:
    """2- and 3-part keys as they might appear in inventory."""
    keys: set[str] = set()
    s, dbs, tgt, dbt, tb, tr = (
        norm_upper(r["Source"]),
        norm_upper(r["DBsource"]),
        norm_upper(r["Target"]),
        norm_upper(r["DBtarget"]),
        norm_upper(r["Tabela"]),
        norm_upper(r["TabRep"]),
    )
    if dbs and tb:
        keys.add(f"{dbs}.{tb}")
    if s and tb:
        keys.add(f"{s}.{tb}")
    if dbt and tb:
        keys.add(f"{dbt}.{tb}")
    if tr and dbt and tb and ("-" in tr or "PDBD" in tr or tr == "BIGQUERY"):
        keys.add(f"{tr}.{dbt}.{tb}")
    if tgt and dbt and tb and tgt not in {"", "BIGQUERY"}:
        keys.add(f"{tgt}.{dbt}.{tb}")
    return keys


def format_repmap_hit(r: dict) -> str:
    return (
        f"Source={r['Source']}|DBsource={r['DBsource']}|Target={r['Target']}|"
        f"DBtarget={r['DBtarget']}|Tabela={r['Tabela']}|TabRep={r['TabRep']}|"
        f"Tipo={r['Tipo']}|Carga={r['Carga']}|row={r['_rownum']}"
    )


def format_replica_primary(r: dict) -> str:
    """Uma linha legível: qual réplica / linha do RepMap bateu."""
    return (
        f"DBsource={r['DBsource']} → DBtarget={r['DBtarget']} | Tabela={r['Tabela']} | "
        f"Target={r['Target']} | TabRep={r['TabRep']} | Tipo={r['Tipo']} | RepDef={r['RepDef']} | row={r['_rownum']}"
    )


def _sybase_collapsed_db_table_pairs(parts: list[str]) -> list[tuple[str, str]]:
    """
    Gera pares (database_or_lib, tabela) para match ao RepMap, colapsando owner intermediário (ex.: DBO).
    parts: segmentos do token já upper e não vazios.
    """
    if not parts:
        return []
    out: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()

    def add(db: str, tbl: str) -> None:
        db, tbl = norm_upper(db), norm_upper(tbl)
        if db and tbl and (db, tbl) not in seen:
            seen.add((db, tbl))
            out.append((db, tbl))

    if len(parts) == 2:
        add(parts[0], parts[1])
        return out

    if len(parts) >= 3 and parts[1] in _SQL_OWNER_SKIP:
        # DBCOR.DBO.TB... → (DBCOR, TB...)
        tbl_rest = ".".join(parts[2:])
        if tbl_rest:
            add(parts[0], tbl_rest)

    # Também tentar último segmento como tabela só se penúltimo for owner skip (casos raros A.X.B com X skip)
    if len(parts) >= 3 and parts[-2] in _SQL_OWNER_SKIP:
        add(parts[0], parts[-1])

    if len(parts) >= 2 and len(out) == 0:
        add(parts[0], parts[-1])

    return out


def _match_schema_table_pair(sch: str, tbl: str, rep_rows: list[dict]) -> list[dict]:
    """sch/tbl upper — mesma lógica que o antigo level 2."""
    hits: list[dict] = []
    for r in rep_rows:
        tbn = norm_upper(r["Tabela"])
        if tbn != tbl:
            continue
        dbs = norm_upper(r["DBsource"])
        src = norm_upper(r["Source"])
        dbt = norm_upper(r["DBtarget"])
        tabrep = norm_upper(r["TabRep"])
        if sch in (dbs, src, dbt, tabrep):
            hits.append(r)
    return hits


def match_artifact(art: str, rep_rows: list[dict], full_key_index: dict[str, list[dict]]) -> tuple[str, str, list[dict]]:
    """
    Returns (match_level, description, matched_rows).
    """
    a = norm_upper(art)
    if not a or a.startswith("/"):
        return "SKIP", "not a qualified table token", []

    parts = [p for p in a.split(".") if p]
    parts_u = [norm_upper(p) for p in parts]

    # Level 1: exact full string against derived keys
    if a in full_key_index:
        hits = full_key_index[a]
        return "EXACT_KEY", f"full token match ({len(hits)} row(s))", hits

    if len(parts_u) >= 3:
        tail3 = ".".join(parts_u[-3:])
        if tail3 in full_key_index:
            hits = full_key_index[tail3]
            return "EXACT_KEY", f"suffix 3-part match ({len(hits)} row(s))", hits
        if a in full_key_index:
            hits = full_key_index[a]
            return "EXACT_KEY", f"full 3+ part ({len(hits)} row(s))", hits

    # Level 2: schema.table (2 partes)
    if len(parts_u) == 2:
        sch, tbl = parts_u[0], parts_u[1]
        hits = _match_schema_table_pair(sch, tbl, rep_rows)
        if len(hits) == 1:
            return "SCHEMA_TABLE", "DBsource|Source|DBtarget|TabRep + Tabela", hits
        if len(hits) > 1:
            return "SCHEMA_TABLE_MULTIPLE", f"{len(hits)} rows share schema/table resolution", hits

    # Level 2b: Sybase/SQL Server database.owner.table — colapsar DBO (etc.) → database.table vs RepMap
    for db, tbl in _sybase_collapsed_db_table_pairs(parts_u):
        key = f"{db}.{tbl}"
        if key in full_key_index:
            hits = full_key_index[key]
            return "SYBASE_OWNER_COLLAPSED", f"match after skipping owner segment (key {key})", hits
        hits = _match_schema_table_pair(db, tbl, rep_rows)
        if len(hits) == 1:
            return "SYBASE_OWNER_COLLAPSED", f"DB+Tabela match with collapsed owner (ex. DBO) → {db}.{tbl}", hits
        if len(hits) > 1:
            return (
                "SYBASE_OWNER_COLLAPSED_MULTIPLE",
                f"{len(hits)} RepMap rows for {db}.{tbl} after owner collapse",
                hits,
            )

    # Level 3: single-part artifact = table name only
    if len(parts_u) == 1:
        tbl = parts_u[0]
        hits = [r for r in rep_rows if norm_upper(r["Tabela"]) == tbl]
        if len(hits) == 1:
            return "TABLE_ONLY_UNIQUE", "unique Tabela in RepMap", hits
        if len(hits) > 1:
            return "TABLE_ONLY_MULTIPLE", f"{len(hits)} rows with same Tabela", hits

    # Level 4: 3-part — try match dataset.table (last two) to DBtarget.Tabela
    if len(parts_u) == 3:
        sch, ds, tbl = parts_u[0], parts_u[1], parts_u[2]
        hits = []
        for r in rep_rows:
            if norm_upper(r["Tabela"]) != tbl:
                continue
            dbt = norm_upper(r["DBtarget"])
            tr = norm_upper(r["TabRep"])
            if dbt == ds and (tr == sch or sch in tr or tr in sch):
                hits.append(r)
            elif dbt == ds:
                hits.append(r)
        if len(hits) >= 1:
            lvl = "THREE_PART" if len(hits) == 1 else "THREE_PART_MULTIPLE"
            return lvl, f"project.dataset.table heuristic ({len(hits)} row(s))", hits

    return "NOT_FOUND", "no RepMap row matched", []


def main():
    root = Path(__file__).resolve().parent
    csv_path = root / "06_prontidao_tabela.csv"
    xlsx_path = root / "src" / "RepMap.xlsx"
    out_path = root / "06_table_in_vs_repmap.csv"

    if len(sys.argv) >= 2:
        csv_path = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        xlsx_path = Path(sys.argv[2])
    if len(sys.argv) >= 4:
        out_path = Path(sys.argv[3])

    rep_rows = load_repmap_rows(xlsx_path)
    full_key_index: dict[str, list[dict]] = {}
    for r in rep_rows:
        for k in repmap_derived_keys(r):
            full_key_index.setdefault(k, []).append(r)

    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        prontidao = list(reader)

    table_in = [row for row in prontidao if norm(row.get("category", "")).lower() == "table_in"]

    fieldnames = [
        "project",
        "area",
        "artifact",
        "category",
        "source_type_06",
        "replica_found",
        "replica_primary",
        "replica_summaries_short",
        "repmap_match_level",
        "repmap_match_description",
        "repmap_hit_count",
        "repmap_matches",
    ]

    with open(out_path, "w", newline="", encoding="utf-8") as out:
        w = csv.DictWriter(out, fieldnames=fieldnames)
        w.writeheader()
        for row in table_in:
            art = norm(row.get("artifact", ""))
            lvl, desc, hits = match_artifact(art, rep_rows, full_key_index)
            has_replica = "YES" if hits and lvl not in ("SKIP",) else "NO"
            primary = format_replica_primary(hits[0]) if hits else ""
            short_list = " | ".join(format_replica_primary(h) for h in hits[:5])
            if len(hits) > 5:
                short_list += f" | ... (+{len(hits) - 5} more)"
            summary = " || ".join(format_repmap_hit(h) for h in hits[:12])
            if len(hits) > 12:
                summary += f" || ... (+{len(hits) - 12} more)"
            w.writerow(
                {
                    "project": row.get("project", ""),
                    "area": row.get("area", ""),
                    "artifact": art,
                    "category": row.get("category", ""),
                    "source_type_06": row.get("source_type", ""),
                    "replica_found": has_replica,
                    "replica_primary": primary,
                    "replica_summaries_short": short_list,
                    "repmap_match_level": lvl,
                    "repmap_match_description": desc,
                    "repmap_hit_count": len(hits),
                    "repmap_matches": summary,
                }
            )

    print(f"Wrote {out_path} ({len(table_in)} table_in rows)")


if __name__ == "__main__":
    main()
